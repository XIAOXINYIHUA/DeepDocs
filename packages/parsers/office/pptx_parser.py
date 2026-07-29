"""PPTX 解析器 —— python-pptx。

支持：
- 幻灯片页面编号
- 标题、文本框内容
- 备注提取
- 表格
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.util import Inches

from ..block import BlockType, ParsedBlock
from ..registry import BaseParser


class PptxParser(BaseParser):
    @property
    def supported_formats(self) -> list[str]:
        return ["pptx"]

    async def validate(self, file_path: str | Path) -> bool:
        try:
            Presentation(str(file_path))
            return True
        except Exception:
            return False

    async def parse(
        self, file_path: str | Path, **kwargs: Any
    ) -> list[Block]:
        prs = Presentation(str(file_path))
        blocks: list[Block] = []

        slide_num = 0
        for slide in prs.slides:
            slide_num += 1

            # 幻灯片标题
            if slide.shapes.title and slide.shapes.title.text.strip():
                title_text = slide.shapes.title.text.strip()
                blocks.append(ParsedBlock(
                      
                    block_type=BlockType.HEADING,
                    page_number=slide_num,
                    heading_path=title_text,
                    heading_level=1,
                    text=title_text,
                    text_length=len(title_text),
                    order_index=len(blocks),
                    depth=1,
                    block_metadata={"slide_number": slide_num, "type": "slide_title"},
                ))

            # 遍历形状
            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue

                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if not text:
                            continue
                        blocks.append(ParsedBlock(
                              
                            block_type=BlockType.PARAGRAPH,
                            page_number=slide_num,
                            heading_path=title_text if slide.shapes.title else None,
                            text=text,
                            text_length=len(text),
                            order_index=len(blocks),
                            depth=0,
                            block_metadata={"slide_number": slide_num},
                        ))

                elif shape.has_table:
                    table = shape.table
                    rows = []
                    for row in table.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        rows.append(cells)

                    if rows:
                        blocks.append(ParsedBlock(
                              
                            block_type=BlockType.TABLE,
                            page_number=slide_num,
                            text=f"[表格: {len(rows)}行]",
                            text_length=0,
                            table_data={
                                "headers": rows[0] if len(rows) > 1 else [],
                                "rows": rows[1:] if len(rows) > 1 else rows,
                                "num_rows": len(rows),
                                "num_cols": len(rows[0]) if rows else 0,
                            },
                            order_index=len(blocks),
                            depth=0,
                            block_metadata={"slide_number": slide_num},
                        ))

            # 备注
            if slide.has_notes_slide:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()
                if notes_text:
                    blocks.append(ParsedBlock(
                          
                        block_type=BlockType.PARAGRAPH,
                        page_number=slide_num,
                        heading_path=title_text if slide.shapes.title else None,
                        text=f"[备注]: {notes_text}",
                        text_length=len(notes_text),
                        order_index=len(blocks),
                        depth=0,
                        block_metadata={"slide_number": slide_num, "type": "notes"},
                    ))

        return blocks
